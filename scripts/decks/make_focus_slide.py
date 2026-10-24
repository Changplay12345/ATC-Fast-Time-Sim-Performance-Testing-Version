# Builds D:\ATC\FTS_Focus_Next.pptx: one white, clean 16:9 slide answering
# "should we finish the software before choosing a server?" for the owner's
# boss. Same style as make_1u_slides.py (Arial, dark text, teal accent).
#
#   pip install python-pptx
#   python scripts/decks/make_focus_slide.py [out.pptx]
#   powershell -File scripts/decks/export_slides.ps1 -Deck <pptx> -OutDir <dir>
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = sys.argv[1] if len(sys.argv) > 1 else r"D:\ATC\FTS_Focus_Next.pptx"
FONT = "Arial"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
RULE = RGBColor(0xD1, 0xD5, 0xDB)
ACCENT = RGBColor(0x0B, 0x72, 0x85)
TINT = RGBColor(0xF3, 0xF6, 0xF9)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
L, W = 0.5, 12.333


def text(x, y, w, h, lines, size=10, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = line if isinstance(line, list) else [(line, bold, color)]
        for t, b, c in runs:
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size), b
            r.font.color.rgb = c
    return tb


def rule(x, y, w, color=RULE, weight=0.75):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    ln.shadow.inherit = False


def heading(x, y, w, t):
    text(x, y, w, 0.3, t, size=11, bold=True, color=ACCENT)
    rule(x, y + 0.32, w)


def bullets(x, y, w, h, items, size=10):
    lines = []
    for lead, rest in items:
        lines.append([("\u2022  " + lead, True, INK), (" " + rest, False, INK)])
    return text(x, y, w, h, lines, size=size, space=5)


# Title and one-line answer
text(L, 0.32, 11.5, 0.5, "What we should focus on next", size=22, bold=True)
text(L, 0.82, 11.8, 0.4, "Finish the software first. Keep the server decision cheap and open until real users have run it.",
     size=13, bold=True, color=ACCENT)

# Three columns
cy, cw, gap = 1.5, 3.95, 0.24
c1, c2, c3 = L, L + cw + gap, L + 2 * (cw + gap)

heading(c1, cy, cw, "Why software first")
bullets(c1, cy + 0.42, cw, 2.6, [
    ("The right server depends on real use,", "and nobody knows that yet. Today's load numbers come from tests, not users. The app has a built-in load meter for later."),
    ("The simulation is cheap to run.", "Nothing about it forces a hardware purchase now."),
    ("The full version will change the numbers.", "Every new feature changes the computing each run needs; sizing done today would be redone."),
    ("Users judge the software, not the server.", ""),
], size=9.5)

heading(c2, cy, cw, "What we measured")
facts = [("10 s", "one 1,000-flight run, one user, on an ordinary 8-core PC"),
         ("< 2 GB", "memory for the whole simulation engine"),
         ("THB 720 / month", "expected cloud cost at the planned use")]
fy = cy + 0.5
for n, l in facts:
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c2), Inches(fy), Inches(cw), Inches(0.72))
    box.fill.solid(); box.fill.fore_color.rgb = TINT; box.line.fill.background(); box.shadow.inherit = False
    text(c2 + 0.12, fy + 0.06, cw - 0.2, 0.34, n, size=17, bold=True, color=ACCENT)
    text(c2 + 0.12, fy + 0.4, cw - 0.2, 0.3, l, size=8.5, color=MUTED)
    fy += 0.84

heading(c3, cy, cw, "Keeping the server decision cheap")
bullets(c3, cy + 0.42, cw, 2.6, [
    ("The desktop app needs no server.", "Testers keep using it whatever we decide later."),
    ("Several people on one PC already works", "over the office network, at no cost."),
    ("A remote demo can run on the cloud trial,", "about THB 720 a month, started in an afternoon, cancelled any month."),
    ("No hardware purchase yet.", "The 1U shortlist and both proposals are ready for when real use justifies it."),
], size=9.5)

# Suggested order: a real sequence, so numbered
oy = 4.55
heading(L, oy, W, "Suggested order")
steps = [
    ("1", "Build the complete version", "using everything learned from the test version."),
    ("2", "Put it in front of real users", "on desktop, plus the cloud trial if a remote demo is wanted."),
    ("3", "Measure for about a month", "read the load meter and the cloud bill."),
    ("4", "Then choose the server", "cloud, rented rack space or an office machine, with evidence."),
]
sw = (W - 3 * 0.2) / 4
sx = L
for n, lead, rest in steps:
    sy = oy + 0.5
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx), Inches(sy), Inches(0.36), Inches(0.36))
    circ.fill.solid(); circ.fill.fore_color.rgb = ACCENT; circ.line.fill.background(); circ.shadow.inherit = False
    tf = circ.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = n; r.font.name, r.font.size, r.font.bold = FONT, Pt(11), True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(sx + 0.46, sy - 0.02, sw - 0.5, 0.3, lead, size=10.5, bold=True)
    text(sx + 0.46, sy + 0.3, sw - 0.5, 0.7, rest, size=9, color=MUTED)
    sx += sw + 0.2

# Trade-off
ty = 6.15
box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(L), Inches(ty), Inches(W), Inches(0.62))
box.fill.solid(); box.fill.fore_color.rgb = TINT; box.line.fill.background(); box.shadow.inherit = False
text(L + 0.12, ty + 0.08, W - 0.3, 0.5,
     [[("The trade-off: ", True, INK), ("the server question stays open for a few months. In return, when it is answered, it is answered with real usage data and the smallest possible spend.", False, INK)]],
     size=10, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT)
